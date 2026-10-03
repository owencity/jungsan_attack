#!/usr/bin/env python3
"""정산어택 DDD 유비쿼터스 언어 워크북을 생성한다.

사용법: python scripts/generate_ddd_domain_language_workbook.py
출력:   docs/ddd-domain-language-workbook.xlsx
"""

from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo


ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "docs" / "ddd-domain-language-workbook.xlsx"

NAVY = "1F4E78"
BLUE = "D9EAF7"
GREEN = "E2F0D9"
YELLOW = "FFF2CC"
RED = "FCE4D6"
GRAY = "E7E6E6"
WHITE = "FFFFFF"
THIN_GRAY = Side(style="thin", color="D9E1F2")


def style_sheet(ws, widths, freeze="A2", filter_ref=None):
    ws.freeze_panes = freeze
    ws.sheet_view.showGridLines = False
    ws.auto_filter.ref = filter_ref or ws.dimensions
    ws.row_dimensions[1].height = 32
    for cell in ws[1]:
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.font = Font(name="맑은 고딕", color=WHITE, bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(bottom=THIN_GRAY)
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.font = Font(name="맑은 고딕", size=10)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = Border(bottom=THIN_GRAY)
    for column, width in widths.items():
        ws.column_dimensions[column].width = width


def add_table(ws, name):
    table = Table(displayName=name, ref=ws.dimensions)
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    ws.add_table(table)


def add_dropdown(ws, cell_range, values):
    validation = DataValidation(
        type="list",
        formula1='"' + ",".join(values) + '"',
        allow_blank=True,
    )
    validation.error = "목록에서 값을 선택해주세요."
    validation.errorTitle = "허용되지 않은 값"
    ws.add_data_validation(validation)
    validation.add(cell_range)


def add_status_colors(ws, column_letter, first_row, last_row):
    rules = {
        "확정": GREEN,
        "검토 필요": YELLOW,
        "초안": BLUE,
        "보류": GRAY,
        "폐기": RED,
    }
    for value, color in rules.items():
        ws.conditional_formatting.add(
            f"{column_letter}{first_row}:{column_letter}{last_row}",
            FormulaRule(
                formula=[f'${column_letter}{first_row}="{value}"'],
                fill=PatternFill("solid", fgColor=color),
            ),
        )


def build_guide(wb):
    ws = wb.active
    ws.title = "사용안내"
    ws.append(["항목", "내용"])
    rows = [
        ("목적", "기획·화면·코드·이벤트에서 같은 단어를 같은 뜻으로 사용하기 위한 DDD 유비쿼터스 언어 워크북"),
        ("권장 순서", "① 용어사전 검토 → ② 상태전이 검토 → ③ 도메인이벤트 검토 → ④ 컨텍스트 경계 검토 → ⑤ 채팅 범위 결정 → ⑥ 의사결정로그 기록"),
        ("작성 원칙 1", "한 용어에는 한 의미만 준다. '정산확인'처럼 여러 상태를 뜻하는 표현은 사용하지 않는다."),
        ("작성 원칙 2", "정의에는 포함되는 것뿐 아니라 포함되지 않는 것과 혼동하면 안 되는 표현도 적는다."),
        ("작성 원칙 3", "명사뿐 아니라 명령, 상태, 불변식, 도메인 이벤트를 함께 정리한다."),
        ("작성 원칙 4", "Bounded Context는 언어와 모델의 경계다. 지금 당장 별도 서비스나 DB로 분리한다는 뜻이 아니다."),
        ("결정 상태", "초안: 최초 제안 / 검토 필요: 선택 필요 / 확정: 팀 언어로 사용 / 보류: 후속 범위 / 폐기: 사용 금지"),
        ("MVP 범위", "포함 / 후속 / 미정 중 선택한다. 기능 범위와 인프라 도입 범위를 섞지 않는다."),
        ("가장 중요한 질문", "채팅방은 모임 전체, 특정 술자리, 이의제기 1:1 중 무엇에 귀속되는가? 복수 종류라면 이름과 권한을 분리한다."),
        ("저장 위치", "이 파일은 생각을 정리하는 초안이다. 확정된 내용은 REQUIREMENTS·ADR·API·코드에 순서대로 반영한다."),
    ]
    for row in rows:
        ws.append(row)
    style_sheet(ws, {"A": 22, "B": 105}, filter_ref=f"A1:B{ws.max_row}")
    add_table(ws, "GuideTable")


def build_glossary(wb):
    ws = wb.create_sheet("용어사전")
    headers = [
        "번호", "Bounded Context", "한글 용어", "Code Term", "분류", "정의",
        "포함", "제외·혼동 금지", "권한·불변식", "예시 문장", "MVP", "결정 상태", "질문·메모",
    ]
    ws.append(headers)
    rows = [
        (1, "회원", "사용자", "User", "Entity", "카카오 로그인을 완료한 서비스 회원", "인증 식별자, 표시 이름", "모임 구성원이나 술자리 참여자와 동일하지 않음", "탈퇴 후에도 과거 정산 표시 식별성은 보존", "사용자가 모임에 가입했다", "포함", "확정", ""),
        (2, "모임", "모임", "Group", "Aggregate", "여러 사용자가 지속적으로 소속되는 공간", "관리자, 구성원, 가입, 차단, 보관", "특정 날짜의 술자리나 정산 자체가 아님", "관리자는 정확히 한 명", "모임에서 새로운 술자리를 만들었다", "포함", "확정", ""),
        (3, "모임", "모임 구성원", "GroupMember", "Entity", "특정 모임에 지속적으로 가입한 사용자", "ACTIVE·LEFT·KICKED 상태", "술자리 참여자와 자동으로 동일하지 않음", "같은 사용자는 모임에 한 행만 존재", "구성원이 모임을 탈퇴했다", "포함", "확정", ""),
        (4, "모임", "모임 관리자", "GroupAdmin", "Role", "모임 관리 권한을 가진 유일한 활성 구성원", "관리자 이전, 강퇴, 차단 해제, 보관", "술자리 생성자·차수 총무와 별개", "모임당 정확히 한 명", "관리자가 구성원을 강퇴했다", "포함", "확정", ""),
        (5, "모임", "모임 차단", "GroupBan", "Entity", "강퇴된 사용자의 모임 재가입을 막는 기록", "차단·해제 주체와 시각", "외부 참여자의 특정 술자리 제외와 다름", "해제 전에는 어떤 가입 경로도 허용하지 않음", "관리자가 모임 차단을 해제했다", "포함", "확정", ""),
        (6, "모임", "모임방", "GroupRoom", "UI/검토", "모임 상세 화면 또는 지속 채팅 공간을 뜻할 수 있는 표현", "미정", "모임과 같은 말로 섞어 쓰지 않음", "도메인 객체인지 화면 이름인지 결정 필요", "사용자가 모임방을 열었다", "미정", "검토 필요", "UI 화면인가, 지속 채팅 공간인가?"),
        (7, "술자리", "술자리", "Gathering", "Aggregate", "특정 날짜에 발생한 하나의 정산 대상 행사", "참여 명단, 차수, 진행 상태", "지속적인 모임과 다름", "groupId는 없을 수도 있음", "술자리 참여자 명단을 마감했다", "포함", "확정", ""),
        (8, "술자리", "술자리 생성자", "GatheringCreator", "Role", "술자리와 초기 차수를 만든 사용자", "차수 생성, 총무 지정, 명단 마감", "돈을 받는 차수 총무와 별개", "생성과 동시에 참여자가 됨", "생성자가 2차 총무를 지정했다", "포함", "확정", ""),
        (9, "술자리", "술자리 참여자", "GatheringParticipant", "Entity", "해당 술자리 정산에 참여하는 로그인 사용자", "구성원 참여자와 외부 참여자", "모임 구성원 자격과 별개", "같은 술자리에 같은 사용자는 한 번만 참여", "참여자가 차수 응답을 제출했다", "포함", "확정", ""),
        (10, "술자리", "외부 참여자", "GuestParticipant", "Role", "모임 구성원은 아니지만 특정 술자리에만 참여한 사용자", "해당 술자리의 응답·결과·송금", "모임 전체 정보 조회 권한 없음", "로그인 사용자여야 함", "외부 참여자가 초대 링크로 참여했다", "포함", "확정", ""),
        (11, "술자리", "참여자 명단 마감", "RosterClosure", "Command/State", "신규 참여를 막고 응답 대상을 고정하는 행위", "마감과 금액 확정 전 재개", "정산 종료와 다름", "금액 확정 후에는 재개 불가", "생성자가 참여자 명단을 마감했다", "포함", "확정", ""),
        (12, "술자리", "차수", "Round", "Entity", "1차·2차처럼 한 명의 결제자가 존재하는 비용 단위", "총액, 주류 항목, 총무, 응답", "술자리 전체와 다름", "술자리 안에서 seq 중복 금지", "2차 비용이 등록되었다", "포함", "확정", ""),
        (13, "술자리", "차수 총무", "RoundManager", "Role", "해당 차수를 실제 결제한 유일한 사람이며 수취인", "비용 입력, 응답 검토, 차수 확정", "모임 관리자·전체 방장과 다름", "해당 술자리의 활성 참여자여야 함", "차수 총무가 1차를 확정했다", "포함", "확정", ""),
        (14, "술자리", "차수 응답", "RoundResponse", "Entity", "참여자가 특정 차수에 제출한 참여·음주 상태", "ABSENT·EXEMPT·SOBER·DRANK", "전역 면제나 단순 참석 boolean이 아님", "활성 참여자는 확정 전 모든 차수에 응답", "참여자가 2차 불참으로 응답했다", "포함", "확정", ""),
        (15, "술자리", "주류 항목", "DrinkItem", "Entity", "술 이름·병 수·병당 가격을 가진 차수 입력", "술값 계산 근거", "택시비 등 기타 항목 제외", "병 수와 단가는 1 이상", "총무가 소주 5병을 입력했다", "포함", "확정", ""),
        (16, "정산", "정산", "SettlementProcess", "Process", "참여 응답 수집부터 금액 확정, 송금 확인, 종료까지의 전체 업무 과정", "계산과 입금 확인 흐름", "술자리나 단순 계산 결과와 같은 말이 아님", "금액 확정 후 계산 입력 변경 금지", "술자리 정산이 진행 중이다", "포함", "초안", "코드의 Settlement aggregate와 명칭 구분 필요"),
        (17, "정산", "정산 계산", "SettlementCalculation", "Domain Service", "차수 입력과 응답으로 부담액·송금액을 계산하는 순수 연산", "Rational 계산, 근거, 송금 목록", "DB 저장이나 입금 상태 변경 제외", "계산은 core 모듈만 수행", "서버가 정산 계산을 요청했다", "포함", "확정", ""),
        (18, "정산", "정산 미리보기", "SettlementPreview", "Value Object", "사용자가 확정 전에 검토하는 계산 결과", "금액, 근거, inputHash, revision", "확정된 송금 증거가 아님", "저장하지 않고 현재 입력으로 계산", "사용자가 정산 미리보기를 확인했다", "포함", "확정", ""),
        (19, "정산", "차수 확정", "RoundConfirmation", "Command/State", "차수 총무가 자기 차수의 비용과 모든 응답을 고정하는 행위", "확정과 금액 확정 전 취소", "전체 금액 확정과 다름", "모든 대상 참여자의 응답 필요", "1차 총무가 차수를 확정했다", "포함", "확정", ""),
        (20, "정산", "금액 확정", "SettlementConfirmation", "Command/State", "모든 계산 입력을 잠그고 송금 명세를 생성하는 행위", "입력 해시 검증, 송금 스냅샷", "입금 확인이나 정산 종료와 다름", "모든 차수 확정과 inputHash 일치 필요", "정산 금액이 확정되었다", "포함", "확정", ""),
        (21, "정산", "확정 정산", "Settlement", "Aggregate", "금액 확정 시 생성되는 계산 버전과 송금 명세의 불변 스냅샷", "inputHash, calculationVersion, grandTotal", "미리보기 전체 breakdown 저장이 아님", "술자리당 MVP 한 건", "확정 정산이 생성되었다", "포함", "확정", ""),
        (22, "송금", "결제", "RoundPayment", "Action", "차수 총무가 음식점 등에 전체 비용을 먼저 지불하는 행위", "차수 원금", "참여자가 총무에게 보내는 송금과 다름", "차수당 결제자는 한 명", "철수가 1차 비용을 결제했다", "포함", "초안", "실제 카드 결제 연동은 범위 밖"),
        (23, "송금", "송금 건", "SettlementTransfer", "Entity", "확정된 송금자→수취인→금액 한 건", "상태, 확인 요청, 확인 시각", "차수 결제나 은행 송금 실행과 다름", "자기 송금 없음, 금액 양수", "민수가 철수에게 12,000원을 송금한다", "포함", "확정", ""),
        (24, "송금", "입금 확인 요청", "PaymentConfirmationRequest", "Command", "송금자가 돈을 보냈다고 수취인에게 알리는 행위", "요청 시각과 송금 상태 변경", "실제 입금 확인과 다름", "송금 당사자만 요청", "민수가 입금 확인을 요청했다", "포함", "확정", ""),
        (25, "송금", "입금 확인", "PaymentConfirmation", "Command/State", "수취인이 실제 입금액을 확인하는 행위", "확인·취소·상태 이력", "송금자의 자기 신고와 다름", "수취인만 확인 가능", "철수가 입금을 확인했다", "포함", "확정", ""),
        (26, "송금", "정산확인", "SettlementCheck", "금지 후보", "차수 확정·금액 확정·입금 확인을 모두 뜻할 수 있는 모호한 표현", "없음", "구체적인 상태 전이 용어로 대체", "API·코드·화면에서 사용하지 않는 것을 권장", "정산확인을 눌렀다", "미정", "검토 필요", "폐기하고 구체 용어로 대체할지 결정"),
        (27, "송금", "이의제기", "Dispute", "Aggregate", "응답 또는 입금 문제를 당사자가 조율하는 절차", "대상, 상태, 당사자, 해결 방식", "일반 채팅과 다름", "열린 이의가 있으면 정산 종료 불가", "수취인이 입금 금액에 이의를 제기했다", "포함", "초안", "응답 이의와 입금 이의를 같은 모델로 둘지 검토"),
        (28, "송금", "정산 종료", "SettlementCompletion", "State", "모든 송금이 확인되고 열린 이의제기가 없는 최종 상태", "완료 시각", "금액 확정과 다름", "종료 후 기존 정산 수정 불가", "모든 입금이 확인되어 정산이 종료되었다", "포함", "확정", ""),
        (29, "대화", "채팅방", "ChatRoom", "Aggregate/검토", "메시지 참여 권한과 귀속 대상을 가진 대화 공간", "모임·술자리·이의제기 중 소유 대상", "소유 대상 없는 범용 방은 지양", "방 종류마다 참여 권한과 수명 정의 필요", "사용자가 채팅방에 입장했다", "미정", "검토 필요", "모임/술자리/이의제기 중 귀속 결정"),
        (30, "대화", "채팅 메시지", "ChatMessage", "Entity", "채팅방에 시간순으로 추가되는 사용자 또는 시스템 메시지", "eventId, sender, type, body, createdAt", "정산 원장 데이터가 아님", "수정·삭제 정책과 중복 eventId 정책 필요", "정산 확정 시스템 메시지가 등록되었다", "미정", "초안", "MongoDB 보존·삭제 정책 필요"),
        (31, "대화", "접속 상태", "Presence", "Ephemeral State", "사용자가 현재 어느 채팅 서버와 방에 연결됐는지 나타내는 휘발 상태", "online set, connection location, TTL", "영구 참여 권한이나 읽음 이력의 원장 아님", "Redis 유실 후 재구성 가능해야 함", "사용자의 접속 상태가 만료되었다", "미정", "초안", "단일 인스턴스 MVP에서 Redis 필요성 검토"),
    ]
    for row in rows:
        ws.append(row)
    style_sheet(
        ws,
        {"A": 7, "B": 16, "C": 22, "D": 30, "E": 18, "F": 48, "G": 35,
         "H": 42, "I": 45, "J": 38, "K": 10, "L": 13, "M": 42},
        freeze="A2",
    )
    add_table(ws, "GlossaryTable")
    add_dropdown(ws, f"K2:K{ws.max_row}", ["포함", "후속", "미정"])
    add_dropdown(ws, f"L2:L{ws.max_row}", ["초안", "검토 필요", "확정", "보류", "폐기"])
    add_status_colors(ws, "L", 2, ws.max_row)


def build_transitions(wb):
    ws = wb.create_sheet("상태전이")
    ws.append(["번호", "대상", "현재 상태", "명령", "실행자", "선행 조건", "다음 상태", "도메인 이벤트", "실패·거부 조건", "결정 상태", "메모"])
    rows = [
        (1, "술자리", "ROSTER_OPEN", "참여자 명단 마감", "술자리 생성자", "신규 참여 처리 완료", "RESPONDING", "RosterClosed", "권한 없음, 이미 금액 확정", "초안", ""),
        (2, "술자리", "RESPONDING", "명단 재개", "술자리 생성자", "금액 확정 전", "ROSTER_OPEN", "RosterReopened", "이미 금액 확정", "초안", ""),
        (3, "차수", "DRAFT", "차수 응답 제출", "응답 본인", "활성 참여자", "DRAFT", "RoundResponseSubmitted", "타인 응답, 차수 확정", "확정", "응답은 상태가 아니라 행 존재로 완료 표현"),
        (4, "차수", "DRAFT", "차수 확정", "차수 총무", "활성 참여자 전원 응답", "CONFIRMED", "RoundConfirmed", "미응답 존재, 총무 아님", "확정", ""),
        (5, "차수", "CONFIRMED", "차수 확정 취소", "차수 총무", "전체 금액 확정 전", "DRAFT", "RoundConfirmationCancelled", "금액 확정 이후", "초안", ""),
        (6, "정산", "ROUND_CONFIRMING", "금액 확정", "권한자", "모든 차수 확정, hash/revision 일치", "AMOUNT_CONFIRMED", "SettlementConfirmed", "입력 변경, 검증 실패", "확정", "확정 송금 스냅샷 생성"),
        (7, "송금 건", "WAITING", "입금 확인 요청", "송금자", "확정된 송금 건", "CONFIRMATION_REQUESTED", "PaymentConfirmationRequested", "당사자 아님", "확정", ""),
        (8, "송금 건", "CONFIRMATION_REQUESTED", "입금 확인", "수취인", "실입금 확인", "CONFIRMED", "PaymentConfirmed", "수취인 아님", "확정", ""),
        (9, "송금 건", "CONFIRMATION_REQUESTED", "이의제기", "수취인", "미입금 또는 금액 불일치", "DISPUTED", "DisputeOpened", "당사자 아님, 열린 이의 중복", "초안", ""),
        (10, "송금 건", "DISPUTED", "조율 종료", "당사자", "합의 또는 오프라인 조율", "WAITING", "DisputeResolved", "종료 권한·합의 조건 미충족", "초안", ""),
        (11, "송금 건", "CONFIRMED", "입금 확인 취소", "수취인", "잘못 확인함", "WAITING", "PaymentConfirmationCancelled", "수취인 아님", "초안", "상태 이력 필수"),
        (12, "정산", "PAYMENT_IN_PROGRESS", "정산 종료", "시스템", "모든 송금 CONFIRMED, 열린 이의 없음", "COMPLETED", "SettlementCompleted", "미확인 송금 또는 열린 이의 존재", "확정", ""),
    ]
    for row in rows:
        ws.append(row)
    style_sheet(ws, {"A": 7, "B": 16, "C": 24, "D": 24, "E": 18, "F": 45, "G": 25, "H": 34, "I": 40, "J": 13, "K": 35})
    add_table(ws, "TransitionTable")
    add_dropdown(ws, f"J2:J{ws.max_row}", ["초안", "검토 필요", "확정", "보류", "폐기"])
    add_status_colors(ws, "J", 2, ws.max_row)


def build_events(wb):
    ws = wb.create_sheet("도메인이벤트")
    ws.append(["번호", "한글 이벤트", "Event Name", "발행 Context", "발생 조건", "필수 데이터", "예상 Consumer", "Outbox", "Kafka 필요 조건", "멱등성 키", "MVP", "결정 상태", "메모"])
    rows = [
        (1, "모임이 생성되었다", "GroupCreated", "모임", "모임 생성 트랜잭션 커밋", "groupId, adminUserId, occurredAt", "알림", "필요", "소비자가 2개 이상일 때", "eventId", "포함", "초안", ""),
        (2, "술자리가 생성되었다", "GatheringCreated", "술자리", "술자리 생성", "gatheringId, groupId?, creatorUserId", "알림, 채팅 시스템 메시지", "필요", "두 consumer를 실제 구현할 때", "eventId", "포함", "초안", "groupId nullable"),
        (3, "참여자 명단이 마감되었다", "RosterClosed", "술자리", "명단 마감 성공", "gatheringId, inputRevision", "알림, 채팅", "필요", "독립 consumer가 필요할 때", "eventId", "포함", "초안", ""),
        (4, "차수 응답이 제출되었다", "RoundResponseSubmitted", "술자리", "본인 응답 저장", "gatheringId, roundId, participantId, responseType", "차수 진행 현황, 알림", "검토", "처리 분리가 필요할 때", "eventId", "포함", "초안", "민감 정보 최소화"),
        (5, "차수가 확정되었다", "RoundConfirmed", "정산", "차수 확정 성공", "gatheringId, roundId, managerParticipantId", "알림, 채팅", "필요", "두 consumer가 독립 재시도할 때", "eventId", "포함", "초안", ""),
        (6, "정산 금액이 확정되었다", "SettlementConfirmed", "정산", "송금 스냅샷 생성", "settlementId, gatheringId, inputHash, calculationVersion", "알림, 채팅, 통계", "필요", "MVP에서 Kafka를 정당화하는 핵심 후보", "settlementId", "포함", "초안", "금액 상세를 이벤트에 과다 노출하지 않음"),
        (7, "입금 확인이 요청되었다", "PaymentConfirmationRequested", "송금", "송금자가 요청", "transferId, senderId, recipientId, requestedAt", "수취인 알림", "필요", "알림 외 consumer가 생길 때", "transferId+상태버전", "포함", "초안", ""),
        (8, "입금이 확인되었다", "PaymentConfirmed", "송금", "수취인이 확인", "transferId, confirmedBy, confirmedAt", "송금자 알림, 완료 판정", "필요", "완료 판정이 비동기여야 할 때", "transferId+상태버전", "포함", "초안", "정산 종료 판정의 동기/비동기 결정 필요"),
        (9, "이의가 제기되었다", "DisputeOpened", "송금", "이의제기 생성", "disputeId, transferId, openedBy", "당사자 알림, 채팅방 생성", "필요", "소비자 분리가 실제 필요할 때", "disputeId", "포함", "초안", ""),
        (10, "정산이 종료되었다", "SettlementCompleted", "정산", "모든 송금 확인, 열린 이의 없음", "settlementId, completedAt", "알림, 통계, 채팅", "필요", "여러 consumer 독립 처리", "settlementId", "포함", "초안", ""),
        (11, "채팅 메시지가 등록되었다", "ChatMessagePosted", "대화", "MongoDB 메시지 저장 성공", "messageId, roomId, senderId, type, createdAt", "WebSocket 전달, 푸시", "별도 설계", "MongoDB↔Kafka 이중 쓰기 실패 정책 확정 후", "messageId", "미정", "검토 필요", "메시지 본문을 Kafka에 포함할지 검토"),
    ]
    for row in rows:
        ws.append(row)
    style_sheet(ws, {"A": 7, "B": 26, "C": 34, "D": 18, "E": 32, "F": 48, "G": 35, "H": 12, "I": 43, "J": 24, "K": 10, "L": 13, "M": 40})
    add_table(ws, "EventTable")
    add_dropdown(ws, f"K2:K{ws.max_row}", ["포함", "후속", "미정"])
    add_dropdown(ws, f"L2:L{ws.max_row}", ["초안", "검토 필요", "확정", "보류", "폐기"])
    add_status_colors(ws, "L", 2, ws.max_row)


def build_contexts(wb):
    ws = wb.create_sheet("컨텍스트맵")
    ws.append(["Bounded Context", "책임", "소유 모델", "주 저장소", "외부에 제공하는 언어", "의존 Context", "통합 방식", "MVP", "결정 상태", "메모"])
    rows = [
        ("회원", "로그인 사용자 식별과 탈퇴", "User", "PostgreSQL", "userId, displayName", "없음", "내부 호출", "포함", "초안", "카카오 연동은 인프라 어댑터"),
        ("모임", "가입·구성원·관리자·차단·보관", "Group, GroupMember, GroupBan", "PostgreSQL", "groupId, membershipStatus", "회원", "내부 호출 + 도메인 이벤트", "포함", "초안", ""),
        ("술자리", "참여 명단·차수·응답·비용 입력", "Gathering, Participant, Round, RoundResponse, DrinkItem", "PostgreSQL", "gatheringId, inputRevision", "회원, 모임", "내부 호출 + Outbox", "포함", "초안", "groupId nullable"),
        ("정산", "순수 계산·미리보기·금액 확정", "SettlementCalculation, Settlement", "PostgreSQL + core", "SettlementConfirmed, TransferSpec", "술자리", "core 함수 호출 + Outbox", "포함", "초안", "계산은 core만 수행"),
        ("송금", "입금 확인 요청·확인·이의·종료 조건", "SettlementTransfer, PaymentHistory, Dispute", "PostgreSQL", "PaymentConfirmed, DisputeOpened", "정산, 회원", "내부 호출 + Kafka 후보", "포함", "초안", "실제 은행 송금은 범위 밖"),
        ("대화", "채팅방·메시지·실시간 전달", "ChatRoom, ChatMessage, Presence", "MongoDB + Redis", "ChatMessagePosted", "회원, 모임/술자리/송금", "Kafka + WebSocket 후보", "미정", "검토 필요", "귀속 대상 먼저 결정"),
        ("알림", "내부 알림과 외부 푸시 재시도", "Notification, NotificationOutbox", "PostgreSQL", "알림 발송 결과", "모든 Context", "Outbox → Kafka 또는 폴러", "포함", "초안", "Kafka 도입 전에도 Outbox 유지"),
    ]
    for row in rows:
        ws.append(row)
    style_sheet(ws, {"A": 18, "B": 40, "C": 50, "D": 25, "E": 38, "F": 24, "G": 36, "H": 10, "I": 13, "J": 38})
    add_table(ws, "ContextTable")
    add_dropdown(ws, f"H2:H{ws.max_row}", ["포함", "후속", "미정"])
    add_dropdown(ws, f"I2:I{ws.max_row}", ["초안", "검토 필요", "확정", "보류", "폐기"])
    add_status_colors(ws, "I", 2, ws.max_row)


def build_chat_decision(wb):
    ws = wb.create_sheet("채팅범위결정")
    ws.append(["후보", "귀속 Aggregate", "참여자", "생성 시점", "종료·보존", "주요 메시지", "MongoDB 키 후보", "Redis presence 키", "장점", "주의점", "MVP 선택", "결정 상태", "메모"])
    rows = [
        ("모임 채팅방", "Group", "활성 모임 구성원", "모임 생성 또는 최초 대화", "모임 보관 후 읽기 전용/정책 필요", "일상 대화, 술자리 공지", "groupId + roomId", "group:{id}:online", "지속적인 관계에 자연스러움", "외부 참여자 접근과 오래된 메시지 보존 정책 필요", "미정", "검토 필요", ""),
        ("술자리 채팅방", "Gathering", "해당 술자리 활성 참여자", "술자리 생성", "정산 종료 후 읽기 전용 또는 보존 기간", "장소, 차수, 비용·정산 시스템 메시지", "gatheringId + roomId", "gathering:{id}:online", "정산 이벤트와 연결이 명확함", "모임 구성원이어도 미참여자는 접근 불가", "미정", "검토 필요", "추천 후보"),
        ("이의제기 대화", "Dispute", "송금 당사자 2명", "이의제기 시작", "해결 후 읽기 전용, 정산 근거로 보존", "입금 불일치 조율", "disputeId", "dispute:{id}:online", "권한과 목적이 가장 명확함", "일반 대화 기능을 대신하지 못함", "미정", "검토 필요", "1:1 대화라는 이름도 고려"),
    ]
    for row in rows:
        ws.append(row)
    style_sheet(ws, {"A": 20, "B": 20, "C": 31, "D": 28, "E": 38, "F": 42, "G": 24, "H": 28, "I": 35, "J": 44, "K": 12, "L": 13, "M": 30})
    add_table(ws, "ChatDecisionTable")
    add_dropdown(ws, f"K2:K{ws.max_row}", ["선택", "제외", "후속", "미정"])
    add_dropdown(ws, f"L2:L{ws.max_row}", ["초안", "검토 필요", "확정", "보류", "폐기"])
    add_status_colors(ws, "L", 2, ws.max_row)


def build_decisions(wb):
    ws = wb.create_sheet("의사결정로그")
    ws.append(["날짜", "주제", "정확히 어떤 문제인가", "고려한 대안", "선택", "선택 이유", "트레이드오프", "나중에 아쉬울 점", "영향 문서·코드", "결정 상태"])
    ws.append(["", "채팅방의 귀속 범위", "모임·술자리·이의제기 중 어떤 대화 공간이 MVP에 필요한가", "① 모임 ② 술자리 ③ 이의제기 ④ 복수 운영", "", "", "", "", "REQUIREMENTS, ADR, MongoDB 키, Redis 키", "검토 필요"])
    ws.append(["", "핵심 DBMS", "PostgreSQL 전환의 실제 이유와 기존 MySQL migration 처리", "① MySQL 유지 ② PostgreSQL 신규 baseline", "", "", "", "", "ADR-011, Liquibase, Compose", "검토 필요"])
    ws.append(["", "Kafka MVP 포함", "독립 consumer와 재시도·이벤트 보존 요구가 실제로 존재하는가", "① Outbox 폴러 ② Outbox→Kafka", "", "", "", "", "ADR-002, 이벤트 계약, Compose", "검토 필요"])
    style_sheet(ws, {"A": 13, "B": 25, "C": 48, "D": 45, "E": 30, "F": 45, "G": 40, "H": 40, "I": 38, "J": 13})
    add_table(ws, "DecisionLogTable")
    add_dropdown(ws, "J2:J200", ["초안", "검토 필요", "확정", "보류", "폐기"])
    add_status_colors(ws, "J", 2, 200)


def main():
    wb = Workbook()
    wb.properties.title = "정산어택 DDD 유비쿼터스 언어 워크북"
    wb.properties.subject = "도메인 언어, 상태 전이, 이벤트, 컨텍스트 경계 검토"
    wb.properties.creator = "jeongsan"

    build_guide(wb)
    build_glossary(wb)
    build_transitions(wb)
    build_events(wb)
    build_contexts(wb)
    build_chat_decision(wb)
    build_decisions(wb)

    for ws in wb.worksheets:
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.sheet_properties.tabColor = NAVY

    OUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT)
    print(f"생성 완료: {OUT}")
    print(f"시트 {len(wb.sheetnames)}개: {', '.join(wb.sheetnames)}")


if __name__ == "__main__":
    main()
